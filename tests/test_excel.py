from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import Category, Product, Unit
from app.services.excel import PRODUCT_HEADERS, export_products, import_products


def make_session():
    db = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(db)
    return sessionmaker(bind=db, autoflush=False, expire_on_commit=False)()


def test_export_products_creates_expected_template(tmp_path):
    session = make_session()
    category = Category(name="Makanan", active=True)
    unit = Unit(name="PCS")
    session.add_all([category, unit])
    session.commit()
    session.add(Product(
        barcode="8990001", name="Biskuit", category_id=category.id, unit_id=unit.id,
        purchase_price=Decimal("5000"), selling_price=Decimal("7000"),
        stock=Decimal("10"), minimum_stock=Decimal("2"), active=True,
    ))
    session.commit()
    path = export_products(session, tmp_path / "produk.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == tuple(PRODUCT_HEADERS)
    assert rows[1][0] == "8990001"
    workbook.close()
    session.close()


def test_import_products_add_and_update_preserves_running_stock(tmp_path):
    session = make_session()
    category = Category(name="Minuman", active=True)
    unit = Unit(name="BOTOL")
    session.add_all([category, unit])
    session.commit()
    path = export_products(session, tmp_path / "template.xlsx")
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.append(["8990002", "Teh", "Minuman", "BOTOL", 2000, 3000, 7, 2, "YA"])
    workbook.save(path)
    result = import_products(session, path, mode="add")
    assert result == {"created": 1, "updated": 0}
    product = session.query(Product).filter_by(barcode="8990002").one()
    assert product.stock == Decimal("7.000")

    product.stock = Decimal("5")
    session.commit()
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(sheet.max_row, 6).value = 3500
    sheet.cell(sheet.max_row, 8).value = 1
    workbook.save(path)
    result = import_products(session, path, mode="update")
    assert result == {"created": 0, "updated": 1}
    session.refresh(product)
    assert product.selling_price == Decimal("3500.00")
    assert product.stock == Decimal("5.000")
    session.close()
