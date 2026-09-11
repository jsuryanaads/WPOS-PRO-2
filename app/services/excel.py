from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ..models import Category, Product, Unit, StockMovement
from .products import create_product, update_product

PRODUCT_HEADERS = [
    "Barcode", "Nama Produk", "Kategori", "Satuan",
    "Harga Beli", "Harga Jual", "Stok Awal", "Stok Minimum", "Aktif",
]


def export_products(session, path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Produk"
    sheet.append(PRODUCT_HEADERS)
    categories = {row.id: row.name for row in session.query(Category).all()}
    units = {row.id: row.name for row in session.query(Unit).all()}
    products = session.query(Product).order_by(Product.name).all()
    for product in products:
        sheet.append([
            product.barcode,
            product.name,
            categories.get(product.category_id, ""),
            units.get(product.unit_id, ""),
            float(product.purchase_price),
            float(product.selling_price),
            float(product.stock),
            float(product.minimum_stock),
            "YA" if product.active else "TIDAK",
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 40)
        sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    return path


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value, default="0"):
    if value is None or _text(value) == "":
        value = default
    return Decimal(str(value))


def _category_id(session, name):
    name = _text(name)
    if not name:
        return None
    row = session.query(Category).filter(Category.name == name).first()
    if row is None:
        raise ValueError(f"Kategori tidak ditemukan: {name}")
    return row.id


def _unit_id(session, name):
    name = _text(name)
    if not name:
        return None
    row = session.query(Unit).filter(Unit.name == name).first()
    if row is None:
        raise ValueError(f"Satuan tidak ditemukan: {name}")
    return row.id


def _active(value):
    text = _text(value).upper()
    if text in ("", "YA", "YES", "TRUE", "1", "AKTIF"):
        return True
    if text in ("TIDAK", "NO", "FALSE", "0", "NONAKTIF"):
        return False
    raise ValueError("Kolom Aktif harus YA atau TIDAK")


def import_products(session, path, mode="add"):
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise ValueError("File Excel kosong")
    headers = [_text(value) for value in rows[0]]
    if headers != PRODUCT_HEADERS:
        raise ValueError("Format Excel tidak sesuai. Gunakan file hasil Export Produk sebagai template.")

    created = updated = 0
    errors = []
    try:
        for row_number, values in enumerate(rows[1:], start=2):
            if not any(value is not None and _text(value) for value in values):
                continue
            try:
                data = dict(zip(PRODUCT_HEADERS, values))
                barcode = _text(data["Barcode"])
                name = _text(data["Nama Produk"])
                if not barcode or not name:
                    raise ValueError("Barcode dan Nama Produk wajib diisi")
                category_id = _category_id(session, data["Kategori"])
                unit_id = _unit_id(session, data["Satuan"])
                existing = session.query(Product).filter_by(barcode=barcode).first()
                if existing:
                    if mode != "update":
                        raise ValueError("Barcode sudah digunakan")
                    update_product(
                        session, existing.id, name=name,
                        purchase_price=_number(data["Harga Beli"]),
                        selling_price=_number(data["Harga Jual"]),
                        minimum_stock=_number(data["Stok Minimum"]),
                        category_id=category_id, unit_id=unit_id,
                        active=_active(data["Aktif"]),
                    )
                    updated += 1
                else:
                    create_product(
                        session, barcode, name,
                        purchase_price=_number(data["Harga Beli"]),
                        selling_price=_number(data["Harga Jual"]),
                        stock=_number(data["Stok Awal"]),
                        minimum_stock=_number(data["Stok Minimum"]),
                        category_id=category_id, unit_id=unit_id,
                    )
                    created += 1
            except Exception as exc:
                errors.append(f"Baris {row_number}: {exc}")
        if errors:
            session.rollback()
            raise ValueError("Import dibatalkan karena ada kesalahan:\n" + "\n".join(errors[:20]))
        return {"created": created, "updated": updated}
    except Exception:
        session.rollback()
        raise
